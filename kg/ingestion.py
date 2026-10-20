"""
ingestion.py — Data Ingestion & Processing Layer (Milestone 1).

Connects to three kinds of enterprise source:

  * Database  — transactions table (CSV export) + SQLite case-management DB
  * Documents — PDF, Word (.docx), Markdown and plain-text files
  * Emails    — RFC-822 .eml files (headers + body, quoted replies stripped)

Every source goes through the same ETL steps:
  Extract  → read raw bytes / rows
  Transform→ clean, normalise, validate, enrich with schema fields
  Load     → handed to the graph builder (kg/pipeline.py) and the RAG index
"""

from __future__ import annotations

import email
import hashlib
import io
import re
import sqlite3
from dataclasses import dataclass, field
from email import policy
from pathlib import Path

import numpy as np
import pandas as pd

from . import schema
from .config import PROJECT_ROOT

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt", ".eml"}


# ═════════════════════════════════════════════════════════════════════════════
# Unstructured sources (documents + emails)
# ═════════════════════════════════════════════════════════════════════════════
@dataclass
class SourceDocument:
    doc_id: str
    title: str
    source_type: str  # pdf | docx | markdown | text | email
    text: str
    path: str = ""
    metadata: dict = field(default_factory=dict)

    @property
    def content_hash(self) -> str:
        return hashlib.sha256(self.text.encode("utf-8")).hexdigest()


_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean_text(text: str) -> str:
    """Normalise whitespace and strip control characters / markdown noise."""
    text = _CONTROL_RE.sub(" ", text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)  # **bold** -> bold
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)  # markdown headings
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _read_email(data: bytes) -> tuple[str, dict]:
    msg = email.message_from_bytes(data, policy=policy.default)
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = body_part.get_content() if body_part else ""
    if body_part is not None and body_part.get_content_type() == "text/html":
        body = re.sub(r"<[^>]+>", " ", body)
    # Drop quoted replies ("> ...") so the same text isn't extracted twice.
    body = "\n".join(line for line in body.splitlines() if not line.lstrip().startswith(">"))
    meta = {
        "from": str(msg.get("From", "")),
        "to": str(msg.get("To", "")),
        "subject": str(msg.get("Subject", "")),
        "date": str(msg.get("Date", "")),
    }
    header = (
        f"Email from {meta['from']} to {meta['to']}.\n"
        f"Subject: {meta['subject']}\nDate: {meta['date']}\n\n"
    )
    return header + body, meta


def parse_file(name: str, data: bytes, path: str = "") -> SourceDocument:
    """Turn raw file bytes into a cleaned SourceDocument."""
    ext = Path(name).suffix.lower()
    meta: dict = {}
    if ext == ".pdf":
        raw, kind = _read_pdf(data), "pdf"
    elif ext == ".docx":
        raw, kind = _read_docx(data), "docx"
    elif ext == ".eml":
        (raw, meta), kind = _read_email(data), "email"
    elif ext in {".md", ".txt"}:
        raw, kind = data.decode("utf-8", errors="replace"), ("markdown" if ext == ".md" else "text")
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    text = clean_text(raw)
    title = meta.get("subject") or Path(name).stem.replace("_", " ")
    doc_id = "doc:" + hashlib.sha1(f"{name}|{text}".encode()).hexdigest()[:12]
    return SourceDocument(doc_id=doc_id, title=title, source_type=kind, text=text,
                          path=path or name, metadata=meta)


def load_documents(folder: Path) -> list[SourceDocument]:
    """Read every supported file under `folder` (recursively)."""
    docs = []
    if not folder.exists():
        return docs
    for p in sorted(folder.rglob("*")):
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                try:
                    rel = p.relative_to(PROJECT_ROOT).as_posix()
                except ValueError:
                    rel = p.name
                docs.append(parse_file(p.name, p.read_bytes(), path=rel))
            except Exception as exc:  # one bad file must not stop the pipeline
                print(f"[ingestion] skipped {p.name}: {exc}")
    # de-duplicate identical content coming from different files
    unique = {}
    for d in docs:
        unique.setdefault(d.content_hash, d)
    return list(unique.values())


# ═════════════════════════════════════════════════════════════════════════════
# Structured source 1 — transactions table
# ═════════════════════════════════════════════════════════════════════════════
REQUIRED_COLUMNS = [
    "transaction_id", "amount", "transaction_hour", "merchant_category",
    "foreign_transaction", "location_mismatch", "device_trust_score",
    "velocity_last_24h", "cardholder_age", "is_fraud",
]

VALIDATION_RULES = {
    "amount >= 0": lambda d: d["amount"] >= 0,
    "transaction_hour in 0..23": lambda d: d["transaction_hour"].between(0, 23),
    "device_trust_score in 0..100": lambda d: d["device_trust_score"].between(0, 100),
    "velocity_last_24h >= 0": lambda d: d["velocity_last_24h"] >= 0,
    "cardholder_age in 18..110": lambda d: d["cardholder_age"].between(18, 110),
    "flags are 0/1": lambda d: d["foreign_transaction"].isin([0, 1])
    & d["location_mismatch"].isin([0, 1]) & d["is_fraud"].isin([0, 1]),
}


@dataclass
class QualityReport:
    source: str
    rows_read: int = 0
    rows_loaded: int = 0
    duplicates_removed: int = 0
    nulls_dropped: int = 0
    rule_failures: dict = field(default_factory=dict)

    def as_dict(self) -> dict:
        return {
            "source": self.source,
            "rows_read": self.rows_read,
            "rows_loaded": self.rows_loaded,
            "duplicates_removed": self.duplicates_removed,
            "nulls_dropped": self.nulls_dropped,
            "rule_failures": self.rule_failures,
        }


def load_transactions(csv_path: Path) -> tuple[pd.DataFrame, QualityReport]:
    """Extract + validate + transform the transactions table."""
    report = QualityReport(source=str(csv_path.name))
    df = pd.read_csv(csv_path)
    report.rows_read = len(df)

    # ── schema check ────────────────────────────────────────────────────────
    df.columns = df.columns.str.strip().str.lower()
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"{csv_path.name} is missing columns: {missing}")

    # ── cleaning ────────────────────────────────────────────────────────────
    before = len(df)
    df = df.drop_duplicates(subset="transaction_id")
    report.duplicates_removed = before - len(df)

    before = len(df)
    df = df.dropna(subset=REQUIRED_COLUMNS)
    report.nulls_dropped = before - len(df)

    df["merchant_category"] = df["merchant_category"].astype(str).str.strip().str.title()

    # ── validation rules ────────────────────────────────────────────────────
    ok = pd.Series(True, index=df.index)
    for name, rule in VALIDATION_RULES.items():
        passed = rule(df)
        failures = int((~passed).sum())
        if failures:
            report.rule_failures[name] = failures
        ok &= passed
    df = df[ok].copy()

    df = transform_transactions(df)
    report.rows_loaded = len(df)
    return df, report


def transform_transactions(df: pd.DataFrame) -> pd.DataFrame:
    """Feature engineering + knowledge-graph enrichment (same logic as the M1 notebook)."""
    df = df.copy()
    int_cols = ["transaction_id", "transaction_hour", "foreign_transaction", "location_mismatch",
                "device_trust_score", "velocity_last_24h", "cardholder_age", "is_fraud"]
    df[int_cols] = df[int_cols].astype(int)

    df["risk_flag"] = df["foreign_transaction"] + df["location_mismatch"]
    for col, new in [("amount", "amount_norm"), ("device_trust_score", "trust_norm"),
                     ("velocity_last_24h", "velocity_norm")]:
        lo, hi = df[col].min(), df[col].max()
        df[new] = 0.0 if hi == lo else (df[col] - lo) / (hi - lo)

    df["tx_key"] = "TX_" + df["transaction_id"].astype(str)
    df["fraud_label"] = np.where(df["is_fraud"] == 1, "Fraud", "Normal")
    df["time_slot"] = df["transaction_hour"].map(schema.time_slot)
    df["customer_segment"] = df["cardholder_age"].map(schema.customer_segment)
    df["risk_factors"] = df.apply(schema.risk_factors_for, axis=1)
    df["risk_factor_count"] = df["risk_factors"].map(len)
    return df


def transaction_to_text(r) -> str:
    """Readable sentence for one transaction (used for RAG + the graph tooltip)."""
    factors = ", ".join(r["risk_factors"]) if r["risk_factors"] else "none"
    return (
        f"Transaction {r['tx_key']} is a {r['fraud_label']} {r['merchant_category']} purchase of "
        f"${r['amount']:.2f} at {int(r['transaction_hour']):02d}:00 ({r['time_slot']}) by a cardholder in "
        f"segment {r['customer_segment']}. Device trust score {int(r['device_trust_score'])}, "
        f"{int(r['velocity_last_24h'])} transactions in the last 24h. Risk factors: {factors}."
    )


# ═════════════════════════════════════════════════════════════════════════════
# Structured source 2 — case-management database
# ═════════════════════════════════════════════════════════════════════════════
def load_case_database(db_path: Path) -> tuple[pd.DataFrame, pd.DataFrame, QualityReport]:
    """Read analysts and fraud cases from the SQLite database."""
    report = QualityReport(source=db_path.name)
    if not db_path.exists():
        return pd.DataFrame(), pd.DataFrame(), report
    with sqlite3.connect(db_path) as con:
        analysts = pd.read_sql_query("SELECT * FROM analysts", con)
        cases = pd.read_sql_query("SELECT * FROM fraud_cases", con)
    report.rows_read = len(cases)
    before = len(cases)
    cases = cases.drop_duplicates(subset="case_id").dropna(subset=["case_id", "transaction_id"])
    report.duplicates_removed = before - len(cases)
    cases["transaction_id"] = cases["transaction_id"].astype(int)
    cases["tx_key"] = "TX_" + cases["transaction_id"].astype(str)
    cases["notes"] = cases["notes"].fillna("").map(clean_text)
    report.rows_loaded = len(cases)
    return analysts, cases, report
